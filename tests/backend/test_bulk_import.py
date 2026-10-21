import io
import sys
import os
import pytest
import openpyxl
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "../../backend")))
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "../..")))

from app.main import app

client = TestClient(app)


@pytest.fixture
def access_token():
    res = client.post("/api/v1/auth/login", json={"username": "access", "password": "12345"})
    return res.json()["access_token"]


def test_download_bulk_template():
    res = client.get("/api/v1/bulk-import/template")
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "student-list" in wb.sheetnames
    sheet = wb["student-list"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == ["Name", "Number", "Address", "Photo File Path"]


def test_bulk_preview_valid_file(access_token):
    # Create valid sample in-memory XLSX
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "student-list"
    ws.append(["Name", "Number", "Address", "Photo File Path"])
    ws.append(["Karthik Raman", "23MBBS901", "Trichy Campus", "photos/karthik.jpg"])
    ws.append(["Deepa Lakshmi", "23MBBS902", "Madurai Road", ""])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    files = {"file": ("student_upload.xlsx", output, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    headers = {"Authorization": f"Bearer {access_token}"}

    res = client.post("/api/v1/bulk-import/preview", files=files, headers=headers)
    assert res.status_code == 200
    data = res.json()
    assert data["total_records"] == 2
    assert data["valid_records"] == 2
    assert data["error_records"] == 0


def test_bulk_preview_invalid_headers(access_token):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "student-list"
    ws.append(["WrongHeader1", "WrongHeader2"])
    ws.append(["John Doe", "12345"])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    files = {"file": ("invalid_upload.xlsx", output, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    headers = {"Authorization": f"Bearer {access_token}"}

    res = client.post("/api/v1/bulk-import/preview", files=files, headers=headers)
    assert res.status_code == 400
    assert "Incorrect column headers" in res.json()["detail"]
